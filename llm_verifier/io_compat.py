"""Workbook / JSON IO helpers + openpyxl capability probe + report printer."""

import json

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    openpyxl = None
    HAS_OPENPYXL = False


def _load_workbook(path):
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is not installed; cannot inspect Excel files")
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def _load_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _iter_sheets_with_rows(wb):
    """Yield (sheet_name, list_of_nonempty_rows) for every worksheet."""
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = list(row)
            if any(v is not None and str(v).strip() != "" for v in vals):
                rows.append(vals)
        yield ws.title, rows


def _emit(report, as_json):
    if as_json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return
    status = "[PASS]" if report.get("pass") else "[FAIL]"
    print(f"{status} {report.get('checker')}")
    lists = ("violations", "mismatches", "reasons")
    for key in lists:
        for v in report.get(key, []):
            if isinstance(v, dict):
                label = v.get("issue") or v.get("kind") or "?"
                print(f"  - {label}: {v.get('detail')}")
            else:
                print(f"  - {v}")
    # warnings shown distinctly, never affect pass/fail
    for v in report.get("warnings", []):
        if isinstance(v, dict):
            print(f"  ~ warning ({v.get('issue')}): {v.get('detail')}")
        else:
            print(f"  ~ warning: {v}")
